"""Workbench Lens — deepest institutional report.

10-sheet XLSX + 4-page PDF covering financial health, risk, customer analytics,
compliance (Indian jurisdiction), and pitch-deck reconciliation.

Usage:
    from reports.workbench_lens import WorkbenchAnalysisResult, WorkbenchLensReport
    xlsx, pdf = WorkbenchLensReport().generate(result)
"""
from __future__ import annotations

import datetime
import logging
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors as rl_colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from analysis.compliance import ComplianceReport
from analysis.customer_analytics import CustomerAnalyticsReport, aggregator_caveat_text
from analysis.financial_health_alerts import classify_nrr
from analysis.financial_analyst import FinancialMetrics
from analysis.reconciliation import ReconciliationReport
from analysis.risk import RiskReport, Severity
from schema.canonical import StatementDocument

log = logging.getLogger(__name__)
_ZERO = Decimal("0")


# ─────────────────────────────────────────────────────────────────────────────
# Input container
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class WorkbenchAnalysisResult:
    doc: StatementDocument
    metrics: FinancialMetrics
    risk_report: RiskReport
    customer_analytics: CustomerAnalyticsReport
    compliance_report: ComplianceReport
    reconciliation_report: ReconciliationReport
    company_name: str | None = None
    narrative: str | None = None  # one-paragraph summary — see analysis/report_narrative.py


# ─────────────────────────────────────────────────────────────────────────────
# XLSX helpers
# ─────────────────────────────────────────────────────────────────────────────

def _inr(v: Decimal | float | None) -> str:
    if v is None:
        return "—"
    f = float(v)
    if f >= 1e7:
        return f"₹{f/1e7:.2f}Cr"
    if f >= 1e5:
        return f"₹{f/1e5:.2f}L"
    return f"₹{f:,.0f}"


_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
_BODY_FONT   = Font(name="Calibri", size=9)
_TITLE_FONT  = Font(name="Calibri", bold=True, size=11)
_NAVY  = PatternFill("solid", fgColor="1F3864")
_SLATE = PatternFill("solid", fgColor="44546A")
_GREEN = PatternFill("solid", fgColor="E2EFDA")
_RED   = PatternFill("solid", fgColor="FCE4D6")
_AMBER = PatternFill("solid", fgColor="FFEB9C")
_THIN  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT   = Alignment(horizontal="left",  vertical="top",    wrap_text=True)


def _header_row(ws, cols: list[str], row: int = 1, fill=_NAVY) -> None:
    for c, val in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.font = _HEADER_FONT
        cell.fill = fill
        cell.border = _THIN
        cell.alignment = _CENTER


def _data_row(ws, values: list, row: int, fills: list | None = None) -> None:
    for c, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.font = _BODY_FONT
        cell.border = _THIN
        cell.alignment = _LEFT
        if fills and c - 1 < len(fills) and fills[c - 1]:
            cell.fill = fills[c - 1]


def _autowidth(ws, min_w: int = 10, max_w: int = 50) -> None:
    for col in ws.columns:
        length = max((len(str(cell.value or "")) for cell in col), default=min_w)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, min_w), max_w)


def _sev_fill(severity: str | Severity) -> PatternFill:
    s = str(severity).upper()
    if "HIGH" in s:
        return _RED
    if "MEDIUM" in s:
        return _AMBER
    return _GREEN


# ─────────────────────────────────────────────────────────────────────────────
# Sheet builders
# ─────────────────────────────────────────────────────────────────────────────

def _sheet_summary(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.active
    ws.title = "Executive Summary"
    m, rr, ca = r.metrics, r.risk_report, r.customer_analytics

    ws.merge_cells("A1:F1")
    ws["A1"] = f"WORKBENCH ANALYSIS — {r.company_name or 'Company'}"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = _NAVY
    ws["A1"].alignment = _CENTER

    # KPI tiles
    kpis = [
        ("Total Revenue", _inr(m.total_revenue)),
        ("Avg Monthly Burn", _inr(m.avg_monthly_burn)),
        ("Runway", f"{float(m.runway_months):.1f} mo" if m.runway_months else "—"),
        ("Risk Flags", str(len(rr.flags))),
        ("Compliance Issues", str(r.compliance_report.high_count) + " HIGH"),
        ("Recon Findings", str(r.reconciliation_report.high_count) + " HIGH"),
    ]
    _header_row(ws, ["Metric", "Value", "Metric", "Value", "Metric", "Value"], row=3)
    ws.cell(row=4, column=1, value=kpis[0][0]).font = _BODY_FONT
    ws.cell(row=4, column=2, value=kpis[0][1]).font = Font(name="Calibri", bold=True, size=9)
    ws.cell(row=4, column=3, value=kpis[1][0]).font = _BODY_FONT
    ws.cell(row=4, column=4, value=kpis[1][1]).font = Font(name="Calibri", bold=True, size=9)
    ws.cell(row=4, column=5, value=kpis[2][0]).font = _BODY_FONT
    ws.cell(row=4, column=6, value=kpis[2][1]).font = Font(name="Calibri", bold=True, size=9)
    ws.cell(row=5, column=1, value=kpis[3][0]).font = _BODY_FONT
    ws.cell(row=5, column=2, value=kpis[3][1]).font = Font(name="Calibri", bold=True, size=9)
    ws.cell(row=5, column=3, value=kpis[4][0]).font = _BODY_FONT
    ws.cell(row=5, column=4, value=kpis[4][1]).font = Font(name="Calibri", bold=True, size=9)
    ws.cell(row=5, column=5, value=kpis[5][0]).font = _BODY_FONT
    ws.cell(row=5, column=6, value=kpis[5][1]).font = Font(name="Calibri", bold=True, size=9)

    row = 7
    if r.narrative:
        ws.merge_cells(f"A{row}:F{row}")
        ws.cell(row=row, column=1, value="SUMMARY").font = Font(name="Calibri", bold=True, color="FFFFFF")
        ws.cell(row=row, column=1).fill = _SLATE
        row += 1
        ws.merge_cells(f"A{row}:F{row}")
        nc = ws.cell(row=row, column=1, value=r.narrative)
        nc.font = _BODY_FONT
        nc.alignment = _LEFT
        ws.row_dimensions[row].height = 45
        row += 2

    # Flag counts by severity — replaces the removed composite risk score
    ws.merge_cells(f"A{row}:F{row}")
    ws.cell(row=row, column=1, value="FLAG SUMMARY").font = Font(name="Calibri", bold=True, color="FFFFFF")
    ws.cell(row=row, column=1).fill = _SLATE
    row += 1
    _header_row(ws, ["Category", "Total", "High", "Medium", "Low"], row=row, fill=_SLATE)
    row += 1
    for label, items in [
        ("Risk flags", rr.flags),
        ("Compliance exceptions", r.compliance_report.exceptions),
        ("Reconciliation findings", r.reconciliation_report.findings),
    ]:
        high = sum(1 for x in items if str(x.severity) == "HIGH")
        med  = sum(1 for x in items if str(x.severity) == "MEDIUM")
        low  = sum(1 for x in items if str(x.severity) == "LOW")
        _data_row(ws, [label, len(items), high, med, low], row=row)
        row += 1
    row += 1

    # Top findings (risk + compliance + reconciliation combined)
    ws.merge_cells(f"A{row}:F{row}")
    ws.cell(row=row, column=1, value="KEY FINDINGS").font = _TITLE_FONT
    ws.cell(row=row, column=1).fill = _SLATE
    ws.cell(row=row, column=1).font = Font(name="Calibri", bold=True, color="FFFFFF")

    row += 1
    _header_row(ws, ["Source", "Severity", "Finding", "Impact", "", ""], row=row, fill=_SLATE)
    row += 1

    for flag in rr.flags[:4]:
        f = _sev_fill(flag.severity)
        _data_row(ws, ["Risk", str(flag.severity), flag.detector_name, flag.description[:100], "", ""], row=row, fills=[f, f])
        row += 1
    for exc in r.compliance_report.exceptions[:3]:
        f = _sev_fill(exc.severity)
        _data_row(ws, ["Compliance", str(exc.severity), exc.rule_name, exc.description[:100], "", ""], row=row, fills=[f, f])
        row += 1
    for finding in r.reconciliation_report.findings[:3]:
        f = _sev_fill(finding.severity)
        _data_row(ws, ["Reconciliation", finding.severity, finding.check_name, finding.description[:100], "", ""], row=row, fills=[f, f])
        row += 1

    _autowidth(ws)


def _sheet_financial(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Financial Health")
    _header_row(ws, ["Month", "Revenue", "Burn", "Net Cash Flow", "Signal"], row=1)
    for i, s in enumerate(r.metrics.monthly_stats, 2):
        net = s.revenue - s.burn
        signal = "▲ Growing" if net > 0 else "▼ Burning"
        fill = _GREEN if net > 0 else _RED
        _data_row(ws, [s.year_month, _inr(s.revenue), _inr(s.burn), _inr(net), signal],
                  row=i, fills=[None, None, None, fill, fill])
    _autowidth(ws)


def _sheet_risk(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Risk Flags")
    _header_row(ws, ["Detector", "Severity", "Score", "Description", "Transaction IDs"], row=1)
    for i, flag in enumerate(r.risk_report.flags, 2):
        ids = ", ".join(flag.triggering_transaction_ids[:5])
        _data_row(ws, [flag.detector_name, str(flag.severity), "", flag.description, ids],
                  row=i, fills=[_sev_fill(flag.severity), _sev_fill(flag.severity)])

    row = len(r.risk_report.flags) + 3
    high = sum(1 for f in r.risk_report.flags if str(f.severity) == "HIGH")
    med  = sum(1 for f in r.risk_report.flags if str(f.severity) == "MEDIUM")
    low  = sum(1 for f in r.risk_report.flags if str(f.severity) == "LOW")
    ws.cell(row=row, column=1, value="Total Flags").font = Font(bold=True)
    ws.cell(row=row, column=2, value=str(len(r.risk_report.flags))).font = Font(bold=True)
    ws.cell(row=row, column=3, value=f"{high} High / {med} Medium / {low} Low").font = Font(bold=True)
    _autowidth(ws)


def _sheet_customers(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Customer Analytics")
    ca = r.customer_analytics

    row = 1

    caveat = aggregator_caveat_text(ca)
    if caveat:
        cell = ws.cell(row=row, column=1, value=caveat)
        cell.font = Font(name="Calibri", bold=True, size=10, color="991B1B")
        row += 2

    ws.cell(row=row, column=1, value="MONTHLY ACTIVE CUSTOMERS + NRR").font = _TITLE_FONT
    row += 1
    _header_row(ws, ["Month", "Active Customers", "New Acquisitions", "Churn Events", "NRR (period-over-period)"], row=row)
    row += 1
    for m in sorted(ca.monthly_active_customers):
        active = ca.monthly_active_customers[m]
        new = len(ca.new_acquisitions.get(m, []))
        churned = sum(1 for e in ca.churn_events if e.last_payment_date.strftime("%Y-%m") <= m)
        nrr = ca.nrr_per_month.get(m)
        nrr_str = f"{nrr*100:.0f}%" if nrr else "—"
        nrr_fill = {"healthy": _GREEN, "below_par": _AMBER, "contraction": _AMBER,
                    "severe": _RED}.get(classify_nrr(nrr)) if nrr is not None else None
        _data_row(ws, [m, active, new, churned, nrr_str], row=row, fills=[None, None, None, None, nrr_fill])
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="CHURNED CUSTOMERS").font = _TITLE_FONT
    row += 1
    _header_row(ws, ["Customer ID", "Last Payment", "Avg Monthly Spend", "Months Active"], row=row)
    row += 1
    for evt in ca.churn_events:
        _data_row(ws, [evt.customer_id, str(evt.last_payment_date),
                       _inr(evt.previous_avg_monthly_spend), evt.months_active_before_churn], row=row)
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="COHORT RETENTION").font = _TITLE_FONT
    row += 1
    if ca.cohort_retention:
        max_offset = max((max(offsets) for offsets in ca.cohort_retention.values() if offsets), default=0)
        header = ["Cohort"] + [f"M+{n}" for n in range(max_offset + 1)]
        _header_row(ws, header, row=row)
        row += 1
        for cohort_m in sorted(ca.cohort_retention):
            ret = ca.cohort_retention[cohort_m]
            values = [cohort_m] + [f"{ret.get(n, 0)*100:.0f}%" if n in ret else "—" for n in range(max_offset + 1)]
            fills = [None] + [(_GREEN if ret.get(n, 0) >= 0.8 else _AMBER if ret.get(n, 0) >= 0.5 else _RED if n in ret else None) for n in range(max_offset + 1)]
            _data_row(ws, values, row=row, fills=fills)
            row += 1

    _autowidth(ws)


def _sheet_compliance(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Compliance Findings")
    _header_row(ws, ["Rule ID", "Rule Name", "Regulatory Citation", "Severity",
                     "Technical Description", "Investor Risk Framing", "Transaction Count"], row=1)
    if not r.compliance_report.exceptions:
        ws.cell(row=2, column=1, value="No compliance exceptions found.").font = _BODY_FONT
        _autowidth(ws)
        return
    for i, exc in enumerate(r.compliance_report.exceptions, 2):
        f = _sev_fill(exc.severity)
        _data_row(ws, [exc.rule_id, exc.rule_name, exc.regulatory_citation,
                       str(exc.severity), exc.description, exc.investor_risk_framing,
                       len(exc.triggering_transaction_ids)], row=i, fills=[None, None, None, f])
    _autowidth(ws)


def _sheet_reconciliation(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Reconciliation")
    if not r.reconciliation_report.claims_provided:
        ws.cell(row=1, column=1, value="No company claims provided. Upload pitch deck or enter claims to enable reconciliation.").font = _BODY_FONT
        _autowidth(ws)
        return
    _header_row(ws, ["Check", "Claimed", "Actual (Bank)", "Direction", "Delta %", "Severity", "Investor Framing"], row=1)
    if not r.reconciliation_report.findings:
        ws.cell(row=2, column=1, value="All claims reconcile within tolerance — no material mismatches found.").font = _BODY_FONT
        _autowidth(ws)
        return
    for i, f in enumerate(r.reconciliation_report.findings, 2):
        sev_fill = _sev_fill(f.severity)
        delta_str = f"{f.delta_pct:+.0f}%" if f.delta_pct is not None else "—"
        _data_row(ws, [f.check_name, f.claimed_value, f.actual_value, str(f.direction),
                       delta_str, f.severity, f.investor_framing], row=i, fills=[None, None, None, None, None, sev_fill])
    _autowidth(ws)


def _sheet_transactions(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("All Transactions")
    _header_row(ws, ["Date", "Description", "Debit", "Credit", "Balance",
                     "Category", "Customer ID", "Related Party", "Anomaly Flags"], row=1)
    for i, t in enumerate(r.doc.transactions, 2):
        flags_str = "; ".join(t.anomaly_flags) if t.anomaly_flags else ""
        rp = t.related_party_match or ("Yes" if t.is_related_party else "")
        fill = _RED if t.is_related_party else ((_AMBER if t.anomaly_flags else None))
        _data_row(ws, [str(t.date), t.description[:80],
                       float(t.debit) if t.debit else "", float(t.credit) if t.credit else "",
                       float(t.balance), str(t.category or ""), t.customer_id or "",
                       rp, flags_str], row=i, fills=[None, None, None, None, None, None, None, fill])
    _autowidth(ws)


def _sheet_related_parties(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Related Parties")
    rp_txns = [t for t in r.doc.transactions if t.is_related_party]
    if not rp_txns:
        ws.cell(row=1, column=1, value="No related-party transactions detected.").font = _BODY_FONT
        _autowidth(ws)
        return
    _header_row(ws, ["Affiliate", "Date", "Description", "Debit", "Credit", "Category"], row=1)
    for i, t in enumerate(rp_txns, 2):
        _data_row(ws, [t.related_party_match or "", str(t.date), t.description[:80],
                       float(t.debit) if t.debit else "", float(t.credit) if t.credit else "",
                       str(t.category or "")], row=i, fills=[_AMBER])
    _autowidth(ws)


def _sheet_customer_master(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    ws = wb.create_sheet("Customer Master")
    from collections import defaultdict

    by_cid: dict[str, dict] = defaultdict(lambda: {
        "first": None, "last": None, "total": _ZERO, "count": 0
    })
    for t in r.doc.transactions:
        if t.customer_id and t.credit:
            d = by_cid[t.customer_id]
            d["total"] += t.credit
            d["count"] += 1
            if d["first"] is None or t.date < d["first"]:
                d["first"] = t.date
            if d["last"] is None or t.date > d["last"]:
                d["last"] = t.date

    _header_row(ws, ["Customer ID", "First Payment", "Last Payment",
                     "Total Revenue", "Payment Count", "Status"], row=1)
    churn_ids = {e.customer_id for e in r.customer_analytics.churn_events}
    latest_month = max(r.customer_analytics.monthly_active_customers, default="")
    # Sourced from CustomerAnalyticsAnalyst's own computed data rather than
    # re-scanning raw transactions — keeps this sheet's "Active" status
    # consistent with what customer_analytics.py actually computed.
    active_ids = r.customer_analytics.active_customer_ids_by_month.get(latest_month, frozenset())

    for i, (cid, data) in enumerate(
        sorted(by_cid.items(), key=lambda x: x[1]["total"], reverse=True), 2
    ):
        status = "Active" if cid in active_ids else ("Churned" if cid in churn_ids else "Inactive")
        fill = _GREEN if status == "Active" else (_RED if status == "Churned" else _AMBER)
        _data_row(ws, [cid, str(data["first"]) if data["first"] else "", str(data["last"]) if data["last"] else "",
                       _inr(data["total"]), data["count"], status], row=i, fills=[None, None, None, None, None, fill])
    _autowidth(ws)


def _sheet_raw_export(wb: Workbook, r: WorkbenchAnalysisResult) -> None:
    import json as _json
    ws = wb.create_sheet("Raw Data Export")
    _header_row(ws, ["transaction_id", "date", "description", "debit", "credit", "balance",
                     "category", "counterparty", "customer_id", "is_related_party",
                     "related_party_match", "anomaly_flags"], row=1)
    for i, t in enumerate(r.doc.transactions, 2):
        _data_row(ws, [
            t.transaction_id, str(t.date), t.description,
            float(t.debit) if t.debit else None, float(t.credit) if t.credit else None,
            float(t.balance),
            str(t.category.value) if t.category else "",
            t.counterparty or "",
            t.customer_id or "",
            t.is_related_party or False,
            t.related_party_match or "",
            _json.dumps(t.anomaly_flags),
        ], row=i)
    _autowidth(ws)


# ─────────────────────────────────────────────────────────────────────────────
# PDF summary
# ─────────────────────────────────────────────────────────────────────────────

def _build_pdf(r: WorkbenchAnalysisResult) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = {
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=14, textColor=rl_colors.HexColor("#1F3864"), spaceAfter=4),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10, textColor=rl_colors.HexColor("#44546A"), spaceAfter=3),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=8, leading=11, spaceAfter=4),
        "label": ParagraphStyle("label", fontName="Helvetica-Bold", fontSize=8),
        "small": ParagraphStyle("small", fontName="Helvetica", fontSize=7, textColor=rl_colors.grey),
        "warning": ParagraphStyle("warning", fontName="Helvetica-Bold", fontSize=8.5,
                                   textColor=rl_colors.HexColor("#991B1B"), spaceAfter=4),
    }

    story = []
    navy = rl_colors.HexColor("#1F3864")
    red = rl_colors.HexColor("#C00000")
    amber = rl_colors.HexColor("#FF8C00")
    green = rl_colors.HexColor("#375623")

    def _sev_color(s: str) -> object:
        if "HIGH" in s.upper():
            return red
        if "MEDIUM" in s.upper():
            return amber
        return green

    # ── Page 1: Financial + Risk ─────────────────────────────────────────────
    story.append(Paragraph(f"WORKBENCH ANALYSIS — {r.company_name or 'Company'}", styles["h1"]))
    story.append(Paragraph(
        f"Period: {r.doc.statement_period_start} → {r.doc.statement_period_end} | "
        f"Generated: {datetime.date.today()}", styles["small"]
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=navy))
    story.append(Spacer(1, 4*mm))

    # KPIs
    m = r.metrics
    kpi_data = [
        ["Total Revenue", _inr(m.total_revenue), "Avg Monthly Burn", _inr(m.avg_monthly_burn)],
        ["Avg Monthly Revenue", _inr(m.avg_monthly_revenue), "Runway", f"{float(m.runway_months):.1f} mo" if m.runway_months else "—"],
        ["Risk Flags", str(len(r.risk_report.flags)),
         "Compliance HIGH", str(r.compliance_report.high_count)],
    ]
    kpi_tbl = Table(kpi_data, colWidths=[40*mm, 40*mm, 40*mm, 40*mm])
    kpi_tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("FONTNAME", (3, 0), (3, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.lightgrey),
        ("BACKGROUND", (0, 0), (0, -1), rl_colors.HexColor("#E9EFF8")),
        ("BACKGROUND", (2, 0), (2, -1), rl_colors.HexColor("#E9EFF8")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(kpi_tbl)
    story.append(Spacer(1, 5*mm))

    if r.narrative:
        story.append(Paragraph("Summary", styles["h2"]))
        story.append(Paragraph(r.narrative, styles["body"]))
        story.append(Spacer(1, 4*mm))

    # Risk flags
    story.append(Paragraph("Risk Flags", styles["h2"]))
    if r.risk_report.flags:
        flag_data = [["Detector", "Severity", "Summary"]]
        for flag in r.risk_report.flags[:6]:
            flag_data.append([flag.detector_name, str(flag.severity), flag.description[:70]])
        ft = Table(flag_data, colWidths=[45*mm, 22*mm, 108*mm])
        ft.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#44546A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, rl_colors.lightgrey),
            ("PADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(ft)
    else:
        story.append(Paragraph("No risk flags detected.", styles["body"]))

    story.append(PageBreak())

    # ── Page 2: Compliance ───────────────────────────────────────────────────
    story.append(Paragraph("Compliance Findings", styles["h1"]))
    story.append(HRFlowable(width="100%", thickness=1, color=navy))
    story.append(Spacer(1, 4*mm))

    if r.compliance_report.exceptions:
        for exc in r.compliance_report.exceptions:
            story.append(Paragraph(f"[{exc.severity}] {exc.rule_name}", styles["h2"]))
            story.append(Paragraph(f"<b>Citation:</b> {exc.regulatory_citation}", styles["body"]))
            story.append(Paragraph(f"<b>Finding:</b> {exc.description}", styles["body"]))
            story.append(Paragraph(f"<i>Investor Impact:</i> {exc.investor_risk_framing}", styles["body"]))
            story.append(HRFlowable(width="100%", thickness=0.5, color=rl_colors.lightgrey))
            story.append(Spacer(1, 2*mm))
    else:
        story.append(Paragraph("No compliance exceptions identified.", styles["body"]))

    story.append(PageBreak())

    # ── Page 3: Reconciliation ───────────────────────────────────────────────
    story.append(Paragraph("Pitch-Deck Reconciliation", styles["h1"]))
    story.append(HRFlowable(width="100%", thickness=1, color=navy))
    story.append(Spacer(1, 4*mm))

    if not r.reconciliation_report.claims_provided:
        story.append(Paragraph("No company claims were provided. Provide pitch-deck metrics to enable reconciliation.", styles["body"]))
    elif not r.reconciliation_report.findings:
        story.append(Paragraph("All provided claims reconcile within tolerance — no material mismatches found.", styles["body"]))
    else:
        rec_data = [["Check", "Claimed", "Actual", "Delta", "Severity"]]
        for f in r.reconciliation_report.findings:
            delta_str = f"{f.delta_pct:+.0f}%" if f.delta_pct is not None else "—"
            rec_data.append([f.check_name, f.claimed_value, f.actual_value, delta_str, f.severity])
        rt = Table(rec_data, colWidths=[50*mm, 30*mm, 30*mm, 22*mm, 22*mm])
        rt.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#44546A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, rl_colors.lightgrey),
            ("PADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(rt)
        story.append(Spacer(1, 4*mm))
        for f in r.reconciliation_report.findings:
            story.append(Paragraph(f"<b>{f.check_name}:</b> {f.investor_framing}", styles["body"]))

    story.append(PageBreak())

    # ── Page 4: Customer Analytics ───────────────────────────────────────────
    story.append(Paragraph("Customer Analytics", styles["h1"]))
    story.append(HRFlowable(width="100%", thickness=1, color=navy))
    story.append(Spacer(1, 4*mm))

    ca = r.customer_analytics
    caveat = aggregator_caveat_text(ca)
    if caveat:
        story.append(Paragraph(caveat, styles["warning"]))
        story.append(Spacer(1, 3*mm))
    if ca.nrr_per_month:
        nrr_data = [["Month", "NRR (period-over-period)"]]
        for mn, nrr in sorted(ca.nrr_per_month.items()):
            nrr_data.append([mn, f"{nrr*100:.0f}%"])
        nt = Table(nrr_data, colWidths=[40*mm, 30*mm])
        nt.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#44546A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, rl_colors.lightgrey),
            ("PADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(Paragraph("Net Revenue Retention (period-over-period)", styles["h2"]))
        story.append(nt)
        story.append(Spacer(1, 4*mm))

    if ca.churn_events:
        story.append(Paragraph(f"Churned Customers: {len(ca.churn_events)}", styles["h2"]))
        churn_data = [["Customer ID", "Last Payment", "Avg Monthly Spend", "Months Active"]]
        for evt in ca.churn_events[:10]:
            churn_data.append([evt.customer_id, str(evt.last_payment_date),
                                _inr(evt.previous_avg_monthly_spend), str(evt.months_active_before_churn)])
        ct = Table(churn_data, colWidths=[50*mm, 30*mm, 40*mm, 30*mm])
        ct.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#C00000")),
            ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, rl_colors.lightgrey),
            ("PADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(ct)

    doc.build(story)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Main report class
# ─────────────────────────────────────────────────────────────────────────────

class WorkbenchLensReport:

    def generate(self, result: WorkbenchAnalysisResult) -> tuple[bytes, bytes]:
        """Returns (xlsx_bytes, pdf_bytes)."""
        # ── XLSX ─────────────────────────────────────────────────────────────
        wb = Workbook()
        _sheet_summary(wb, result)
        _sheet_financial(wb, result)
        _sheet_risk(wb, result)
        _sheet_customers(wb, result)
        _sheet_compliance(wb, result)
        _sheet_reconciliation(wb, result)
        _sheet_transactions(wb, result)
        _sheet_related_parties(wb, result)
        _sheet_customer_master(wb, result)
        _sheet_raw_export(wb, result)

        buf = BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        # ── PDF ──────────────────────────────────────────────────────────────
        pdf_bytes = _build_pdf(result)

        return xlsx_bytes, pdf_bytes

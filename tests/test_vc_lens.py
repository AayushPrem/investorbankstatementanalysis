"""Tests for reports/vc_lens.py — VCLensReport."""
from __future__ import annotations

import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from io import BytesIO

from analysis.customer_analytics import (
    ChurnEvent,
    ConcentrationPoint,
    CustomerAnalyticsReport,
)
from analysis.financial_analyst import FinancialMetrics, MonthlyStats
from analysis.risk import Flag, RiskReport, Severity
from reports.vc_lens import AnalysisResult, VCLensReport
from schema.canonical import (
    CanonicalTransaction,
    SourceReference,
    StatementDocument,
    TransactionCategory,
)

# ---------------------------------------------------------------------------
# Test fixtures / builders
# ---------------------------------------------------------------------------

_SRC = SourceReference(file_path="t.pdf", page=0, row=0, raw_text="")


def _txn(tid: str, credit: int | None = None, debit: int | None = None,
         category: TransactionCategory | None = TransactionCategory.REVENUE,
         is_related_party: bool = False,
         related_party_match: str | None = None,
         customer_id: str | None = None,
         month: int = 1) -> CanonicalTransaction:
    d = datetime.date(2024, month, 15)
    if debit is not None:
        return CanonicalTransaction(
            transaction_id=tid, date=d, description=tid,
            debit=Decimal(str(debit)), balance=Decimal("500000"),
            source_reference=_SRC, category=category,
            is_related_party=is_related_party,
            related_party_match=related_party_match,
            customer_id=customer_id,
        )
    return CanonicalTransaction(
        transaction_id=tid, date=d, description=tid,
        credit=Decimal(str(credit or 50000)), balance=Decimal("600000"),
        source_reference=_SRC, category=category,
        is_related_party=is_related_party,
        related_party_match=related_party_match,
        customer_id=customer_id,
    )


def _doc(*txns: CanonicalTransaction) -> StatementDocument:
    return StatementDocument(
        account_id="ACC001",
        statement_period_start=datetime.date(2024, 1, 1),
        statement_period_end=datetime.date(2024, 6, 30),
        source_format="test",
        bank_name="Test Bank",
        transactions=list(txns),
    )


def _metrics(monthly_stats: list[MonthlyStats] | None = None) -> FinancialMetrics:
    stats = monthly_stats or [
        MonthlyStats("2024-01", Decimal("200000"), Decimal("100000")),
        MonthlyStats("2024-02", Decimal("220000"), Decimal("110000")),
        MonthlyStats("2024-03", Decimal("240000"), Decimal("120000")),
    ]
    total_rev  = sum(m.revenue for m in stats)
    total_burn = sum(m.burn    for m in stats)
    n = len(stats)
    avg_rev  = total_rev  / n
    avg_burn = total_burn / n
    return FinancialMetrics(
        period_months=n,
        total_revenue=total_rev,
        total_burn=total_burn,
        closing_balance=Decimal("300000"),
        monthly_stats=stats,
        avg_monthly_revenue=avg_rev,
        avg_monthly_burn=avg_burn,
        runway_months=Decimal("3"),
        mom_revenue_growth_rates=[Decimal("0.10"), Decimal("0.09")],
        avg_mom_growth=Decimal("0.095"),
        revenue_hhi=Decimal("0.3"),
        top_customer_revenue_pct=Decimal("0.4"),
    )


def _risk_report(flags: list[Flag] | None = None) -> RiskReport:
    return RiskReport(
        flags=flags or [],
        composite_score=0.0,
        narrative="No flags.",
    )


def _customer_analytics() -> CustomerAnalyticsReport:
    return CustomerAnalyticsReport(
        monthly_active_customers={"2024-01": 5, "2024-02": 6, "2024-03": 7},
        monthly_active_trend=1.0,
        churn_events=[
            ChurnEvent(
                customer_id="cust_abc",
                last_payment_date=datetime.date(2024, 1, 15),
                previous_avg_monthly_spend=Decimal("50000"),
                months_active_before_churn=3,
            )
        ],
        new_acquisitions={"2024-01": ["cust_abc", "cust_def"]},
        nrr_per_month={"2024-02": 1.05, "2024-03": 0.95},
        cohort_retention={"2024-01": {0: 1.0, 1: 0.8, 2: 0.6}},
        concentration_trajectory=[
            ConcentrationPoint(month="2024-01", top3_share=0.8, top10_share=1.0),
            ConcentrationPoint(month="2024-02", top3_share=0.7, top10_share=1.0),
        ],
        payment_regularity_alerts=[],
    )


def _sample_result(
    flags: list[Flag] | None = None,
    related_party: bool = False,
    company_name: str | None = "Acme Pvt Ltd",
) -> AnalysisResult:
    txns = [
        _txn("t1", credit=200000, customer_id="cust_A", month=1),
        _txn("t2", debit=100000, category=TransactionCategory.VENDOR_PAYMENT, month=1),
        _txn("t3", credit=220000, customer_id="cust_B", month=2),
    ]
    if related_party:
        txns.append(_txn("t4", credit=50000, category=TransactionCategory.REVENUE,
                         is_related_party=True, related_party_match="Acme Holdings",
                         customer_id="cust_rp", month=2))
    return AnalysisResult(
        doc=_doc(*txns),
        metrics=_metrics(),
        risk_report=_risk_report(flags),
        customer_analytics=_customer_analytics(),
        company_name=company_name,
    )


def _load_wb(xlsx_bytes: bytes):
    return load_workbook(BytesIO(xlsx_bytes))


# ---------------------------------------------------------------------------
# TestVCLensReportGeneration
# ---------------------------------------------------------------------------

class TestVCLensReportGeneration:
    def test_returns_tuple_of_two_byte_strings(self) -> None:
        result = _sample_result()
        xlsx, pdf = VCLensReport().generate(result)
        assert isinstance(xlsx, bytes)
        assert isinstance(pdf, bytes)

    def test_xlsx_bytes_non_empty(self) -> None:
        xlsx, _ = VCLensReport().generate(_sample_result())
        assert len(xlsx) > 1000

    def test_pdf_bytes_non_empty(self) -> None:
        _, pdf = VCLensReport().generate(_sample_result())
        assert len(pdf) > 1000

    def test_pdf_starts_with_pdf_magic_bytes(self) -> None:
        _, pdf = VCLensReport().generate(_sample_result())
        assert pdf[:4] == b"%PDF"

    def test_xlsx_is_valid_workbook(self) -> None:
        xlsx, _ = VCLensReport().generate(_sample_result())
        wb = _load_wb(xlsx)
        assert wb is not None


# ---------------------------------------------------------------------------
# TestSheetPresence
# ---------------------------------------------------------------------------

class TestSheetPresence:
    def setup_method(self) -> None:
        xlsx, _ = VCLensReport().generate(_sample_result())
        self.wb = _load_wb(xlsx)

    def test_exactly_six_sheets(self) -> None:
        assert len(self.wb.sheetnames) == 6

    def test_sheet1_summary_present(self) -> None:
        assert "Summary" in self.wb.sheetnames

    def test_sheet2_financial_health_present(self) -> None:
        assert "Financial Health" in self.wb.sheetnames

    def test_sheet3_red_flags_present(self) -> None:
        assert "Red Flags" in self.wb.sheetnames

    def test_sheet4_customer_analytics_present(self) -> None:
        assert "Customer Analytics" in self.wb.sheetnames

    def test_sheet5_all_transactions_present(self) -> None:
        assert "All Transactions" in self.wb.sheetnames

    def test_sheet6_related_parties_present(self) -> None:
        assert "Related Parties" in self.wb.sheetnames

    def test_sheets_in_correct_order(self) -> None:
        expected = [
            "Summary", "Financial Health", "Red Flags",
            "Customer Analytics", "All Transactions", "Related Parties",
        ]
        assert self.wb.sheetnames == expected


# ---------------------------------------------------------------------------
# TestSummarySheetContent
# ---------------------------------------------------------------------------

class TestSummarySheetContent:
    def _get_all_cell_values(self, wb, sheet_name: str) -> dict[str, object]:
        ws = wb[sheet_name]
        return {f"{ws.cell(r, 1).value}": ws.cell(r, 2).value
                for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value is not None}

    def test_total_revenue_matches_metrics(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]

        # Scan all rows to find "Total Revenue" label
        found = None
        for row in ws.iter_rows():
            if row[0].value == "Total Revenue":
                found = row[1].value
                break
        assert found is not None
        assert abs(found - float(result.metrics.total_revenue)) < 0.01

    def test_total_burn_matches_metrics(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]
        found = None
        for row in ws.iter_rows():
            if row[0].value == "Total Burn":
                found = row[1].value
                break
        assert found is not None
        assert abs(found - float(result.metrics.total_burn)) < 0.01

    def test_risk_score_matches_report(self) -> None:
        flags = [Flag(
            detector_name="structuring",
            severity=Severity.HIGH,
            triggering_transaction_ids=["t1"],
            description="Test flag",
        )]
        result = _sample_result(flags=flags)
        result.risk_report.composite_score = 42.0
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]
        found = None
        for row in ws.iter_rows():
            if row[0].value == "Composite Risk Score (/100)":
                found = row[1].value
                break
        assert found is not None
        assert abs(found - 42.0) < 0.01

    def test_summary_contains_company_name(self) -> None:
        result = _sample_result(company_name="Acme Pvt Ltd")
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]
        cell_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any("Acme Pvt Ltd" in str(v) for v in cell_values if v)

    def test_summary_has_top5_flags_section(self) -> None:
        flags = [
            Flag(detector_name=f"det_{i}", severity=Severity.MEDIUM,
                 triggering_transaction_ids=["t1"], description=f"Flag {i}")
            for i in range(7)
        ]
        result = _sample_result(flags=flags)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]
        found = any("FINDINGS" in str(ws.cell(r, 1).value or "").upper()
                    for r in range(1, ws.max_row + 1))
        assert found


# ---------------------------------------------------------------------------
# TestFinancialHealthSheet
# ---------------------------------------------------------------------------

class TestFinancialHealthSheet:
    def test_one_row_per_month(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Financial Health"]
        # Row 1 = header, row 2 = explanatory note, data starts row 3
        # Count rows with numeric revenue value in col 2
        data_rows = sum(1 for r in range(3, ws.max_row + 1)
                        if isinstance(ws.cell(r, 2).value, (int, float)))
        assert data_rows == len(result.metrics.monthly_stats)

    def test_header_row_has_month_column(self) -> None:
        xlsx, _ = VCLensReport().generate(_sample_result())
        wb = _load_wb(xlsx)
        ws = wb["Financial Health"]
        header = [ws.cell(1, c).value for c in range(1, 6)]
        assert "Month" in header

    def test_revenue_values_are_numeric(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Financial Health"]
        # Data rows start at row 3 (row 2 is an explanatory note)
        for r in range(3, ws.max_row + 1):
            if ws.cell(r, 1).value is None:
                break
            assert isinstance(ws.cell(r, 2).value, (int, float))

    def test_revenue_values_match_monthly_stats(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Financial Health"]
        # Row 1 = header, row 2 = explanatory note row, data starts at row 3
        for i, ms in enumerate(result.metrics.monthly_stats, 3):
            cell_val = ws.cell(i, 2).value
            assert cell_val is not None
            assert abs(cell_val - float(ms.revenue)) < 0.01


# ---------------------------------------------------------------------------
# TestRedFlagsSheet
# ---------------------------------------------------------------------------

class TestRedFlagsSheet:
    def test_one_row_per_flag(self) -> None:
        flags = [
            Flag(detector_name="structuring", severity=Severity.HIGH,
                 triggering_transaction_ids=["t1"], description="Structuring detected"),
            Flag(detector_name="round_tripping", severity=Severity.MEDIUM,
                 triggering_transaction_ids=["t2", "t3"], description="Round trip"),
        ]
        result = _sample_result(flags=flags)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Red Flags"]
        # Count rows that have a severity value in col 2 (HIGH/MEDIUM/LOW)
        flag_rows = sum(1 for r in range(2, ws.max_row + 1)
                        if ws.cell(r, 2).value in ("HIGH", "MEDIUM", "LOW"))
        assert flag_rows == 2

    def test_no_flags_sheet_has_no_data_rows(self) -> None:
        result = _sample_result(flags=[])
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Red Flags"]
        # No severity cells should exist when there are no flags
        flag_rows = sum(1 for r in range(2, ws.max_row + 1)
                        if ws.cell(r, 2).value in ("HIGH", "MEDIUM", "LOW"))
        assert flag_rows == 0

    def test_severity_column_present(self) -> None:
        flags = [Flag(detector_name="structuring", severity=Severity.HIGH,
                      triggering_transaction_ids=["t1"], description="Test")]
        result = _sample_result(flags=flags)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Red Flags"]
        assert ws.cell(2, 2).value == "HIGH"

    def test_detector_name_in_first_column(self) -> None:
        flags = [Flag(detector_name="my_detector", severity=Severity.LOW,
                      triggering_transaction_ids=[], description="desc")]
        result = _sample_result(flags=flags)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Red Flags"]
        assert ws.cell(2, 1).value == "my_detector"

    def test_triggering_ids_in_fifth_column(self) -> None:
        flags = [Flag(detector_name="det", severity=Severity.LOW,
                      triggering_transaction_ids=["txn_A", "txn_B"], description="d")]
        result = _sample_result(flags=flags)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Red Flags"]
        # Column 4 is now "What This Means", column 5 is triggering IDs
        ids_cell = ws.cell(2, 5).value or ""
        assert "txn_A" in ids_cell
        assert "txn_B" in ids_cell


# ---------------------------------------------------------------------------
# TestCustomerAnalyticsSheet
# ---------------------------------------------------------------------------

class TestCustomerAnalyticsSheet:
    def test_active_customer_data_present(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Customer Analytics"]
        # Find a cell containing a month label from active_customers
        all_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any("Jan 2024" in str(v) for v in all_values if v)

    def test_nrr_data_present(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Customer Analytics"]
        # NRR section should have some numeric values in column 2
        numeric_vals = [
            ws.cell(r, 2).value for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(r, 2).value, (int, float))
        ]
        assert len(numeric_vals) > 0

    def test_churn_event_customer_id_present(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Customer Analytics"]
        all_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any("cust_abc" in str(v) for v in all_values if v)

    def test_cohort_retention_section_present(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Customer Analytics"]
        all_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any("COHORT" in str(v).upper() for v in all_values if v)

    def test_concentration_trajectory_section_present(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Customer Analytics"]
        all_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any("CONCENTRATION" in str(v).upper() for v in all_values if v)


# ---------------------------------------------------------------------------
# TestAllTransactionsSheet
# ---------------------------------------------------------------------------

class TestAllTransactionsSheet:
    def test_all_transactions_rows_present(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["All Transactions"]
        data_rows = sum(1 for row in ws.iter_rows(min_row=2, values_only=True)
                        if row[0] is not None)
        assert data_rows == len(result.doc.transactions)

    def test_header_includes_id_date_description(self) -> None:
        xlsx, _ = VCLensReport().generate(_sample_result())
        wb = _load_wb(xlsx)
        ws = wb["All Transactions"]
        header = [ws.cell(1, c).value for c in range(1, 11)]
        assert "ID" in header
        assert "Date" in header
        assert "Description" in header

    def test_transaction_ids_appear_in_rows(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["All Transactions"]
        ids_in_sheet = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
        expected_ids = {t.transaction_id for t in result.doc.transactions}
        assert expected_ids == ids_in_sheet

    def test_category_column_populated(self) -> None:
        result = _sample_result()
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["All Transactions"]
        categories = [ws.cell(r, 7).value for r in range(2, ws.max_row + 1)
                      if ws.cell(r, 1).value]
        assert all(c is not None and c != "" for c in categories)


# ---------------------------------------------------------------------------
# TestRelatedPartiesSheet
# ---------------------------------------------------------------------------

class TestRelatedPartiesSheet:
    def test_related_party_transactions_present(self) -> None:
        result = _sample_result(related_party=True)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Related Parties"]
        data_rows = sum(1 for row in ws.iter_rows(min_row=2, values_only=True)
                        if row[0] is not None and "t4" in str(row[0]))
        assert data_rows == 1

    def test_no_related_party_transactions_means_no_data_rows(self) -> None:
        result = _sample_result(related_party=False)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Related Parties"]
        # Only header row (row 1) should be populated in the main section
        data_rows = sum(1 for row in ws.iter_rows(min_row=2, max_row=2, values_only=True)
                        if row[0] is not None and str(row[0]).startswith("t"))
        assert data_rows == 0

    def test_affiliate_column_populated(self) -> None:
        result = _sample_result(related_party=True)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Related Parties"]
        # Find row with t4
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "t4":
                assert row[5].value == "Acme Holdings"
                break

    def test_aggregate_section_present(self) -> None:
        result = _sample_result(related_party=True)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Related Parties"]
        all_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any("AGGREGATED" in str(v).upper() for v in all_values if v)


# ---------------------------------------------------------------------------
# TestEdgeCases
# ---------------------------------------------------------------------------

class TestEdgeCases:
    def test_no_transactions_generates_without_error(self) -> None:
        result = AnalysisResult(
            doc=_doc(),
            metrics=_metrics(),
            risk_report=_risk_report(),
            customer_analytics=CustomerAnalyticsReport(
                monthly_active_customers={},
                monthly_active_trend=0.0,
                churn_events=[],
                new_acquisitions={},
                nrr_per_month={},
                cohort_retention={},
                concentration_trajectory=[],
                payment_regularity_alerts=[],
            ),
        )
        xlsx, pdf = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        assert len(wb.sheetnames) == 6

    def test_many_flags_capped_in_summary(self) -> None:
        flags = [
            Flag(detector_name=f"det_{i}", severity=Severity.LOW,
                 triggering_transaction_ids=[], description=f"Flag {i}")
            for i in range(10)
        ]
        result = _sample_result(flags=flags)
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]
        # Count detector name rows in summary (those starting with "[")
        flag_rows = sum(1 for r in range(1, ws.max_row + 1)
                        if str(ws.cell(r, 1).value or "").startswith("["))
        assert flag_rows <= 6  # summary shows up to 6 top flags

    def test_no_company_name_uses_account_id(self) -> None:
        result = _sample_result(company_name=None)
        result.company_name = None
        xlsx, _ = VCLensReport().generate(result)
        wb = _load_wb(xlsx)
        ws = wb["Summary"]
        all_vals = [ws.cell(r, 1).value for r in range(1, min(5, ws.max_row) + 1)]
        assert any("ACC001" in str(v) for v in all_vals if v)

    def test_profitable_company_shows_no_runway_number(self) -> None:
        result = _sample_result()
        result.metrics = _metrics()
        result.metrics.runway_months = None  # profitable — no runway needed
        xlsx, pdf = VCLensReport().generate(result)
        assert len(xlsx) > 0 and len(pdf) > 0

    def test_generate_is_idempotent(self) -> None:
        result = _sample_result()
        reporter = VCLensReport()
        xlsx1, pdf1 = reporter.generate(result)
        xlsx2, pdf2 = reporter.generate(result)
        # Both outputs should be valid and same size (deterministic)
        assert len(xlsx1) == len(xlsx2)

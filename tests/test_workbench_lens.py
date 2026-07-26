"""Tests for reports/workbench_lens.py — WorkbenchLensReport.

Wave 4.1: this file previously did not exist at all — the Workbench lens
(the deepest, 10-sheet + 4-page report, aimed at in-house finance teams) had
zero test coverage. These tests verify that specific KNOWN values from the
input fixtures actually render correctly in both the XLSX sheets and the PDF
pages — not just that bytes were produced.
"""
from __future__ import annotations

import datetime
import json
from decimal import Decimal
from io import BytesIO

import pdfplumber
import pytest
from openpyxl import load_workbook

from analysis.compliance import ComplianceReport
from analysis.customer_analytics import (
    ChurnEvent,
    ConcentrationPoint,
    CustomerAnalyticsReport,
)
from analysis.financial_analyst import FinancialMetrics, MonthlyStats
from analysis.reconciliation import MismatchDirection, ReconciliationFinding, ReconciliationReport
from analysis.risk import Flag, RiskReport, Severity
from jurisdictions import ComplianceException
from reports.workbench_lens import WorkbenchAnalysisResult, WorkbenchLensReport
from schema.canonical import (
    CanonicalTransaction,
    SourceReference,
    StatementDocument,
    TransactionCategory,
)

_SRC = SourceReference(file_path="t.pdf", page=0, row=0, raw_text="")


def _txn(
    tid: str, *, date: datetime.date, credit: int | None = None, debit: int | None = None,
    category: TransactionCategory | None = TransactionCategory.REVENUE,
    customer_id: str | None = None, is_related_party: bool = False,
    related_party_match: str | None = None, anomaly_flags: list[str] | None = None,
) -> CanonicalTransaction:
    kwargs = dict(
        transaction_id=tid, date=date, description=f"desc-{tid}",
        balance=Decimal("500000"), source_reference=_SRC, category=category,
        customer_id=customer_id, is_related_party=is_related_party,
        related_party_match=related_party_match, anomaly_flags=anomaly_flags or [],
    )
    if debit is not None:
        kwargs["debit"] = Decimal(str(debit))
    else:
        kwargs["credit"] = Decimal(str(credit or 0))
    return CanonicalTransaction(**kwargs)


def _doc() -> StatementDocument:
    return StatementDocument(
        account_id="ACC777", statement_period_start=datetime.date(2024, 1, 1),
        statement_period_end=datetime.date(2024, 2, 29), source_format="test",
        bank_name="HDFC",
        transactions=[
            _txn("t1", date=datetime.date(2024, 1, 15), credit=300000, customer_id="cust_alpha"),
            _txn("t2", date=datetime.date(2024, 2, 15), credit=320000, customer_id="cust_alpha"),
            _txn("t3", date=datetime.date(2024, 1, 20), debit=150000, category=TransactionCategory.VENDOR_PAYMENT),
            _txn("t4", date=datetime.date(2024, 2, 5), debit=75000, category=TransactionCategory.VENDOR_PAYMENT,
                 is_related_party=True, related_party_match="Founder Family Ventures Pvt Ltd"),
            _txn("t5", date=datetime.date(2024, 1, 10), debit=20000, category=TransactionCategory.FEES,
                 anomaly_flags=["round_amount"]),
        ],
    )


def _metrics() -> FinancialMetrics:
    stats = [
        MonthlyStats("2024-01", Decimal("300000"), Decimal("170000")),
        MonthlyStats("2024-02", Decimal("320000"), Decimal("75000")),
    ]
    return FinancialMetrics(
        period_months=2, total_revenue=Decimal("620000"), total_burn=Decimal("245000"),
        closing_balance=Decimal("2000000"), monthly_stats=stats,
        avg_monthly_revenue=Decimal("310000"), avg_monthly_burn=Decimal("122500"),
        runway_months=Decimal("16.3"), mom_revenue_growth_rates=[Decimal("0.0667")],
        avg_mom_growth=Decimal("0.0667"), revenue_hhi=0.5, top_customer_revenue_pct=1.0,
    )


def _risk_report() -> RiskReport:
    return RiskReport(
        flags=[
            Flag(detector_name="founder_over_extraction", severity=Severity.HIGH,
                 triggering_transaction_ids=["t4"],
                 description="Founder withdrawals of Rs.75000 exceed declared salary threshold"),
        ],
        composite_score=42.0,
        narrative="One HIGH-severity flag: founder over-extraction.",
    )


def _compliance_report() -> ComplianceReport:
    return ComplianceReport(
        exceptions=[
            ComplianceException(
                rule_id="IND_RPT_CONCENTRATION", rule_name="Related-Party Transaction Concentration",
                regulatory_citation="Companies Act §188", severity=Severity.HIGH,
                description="Related-party outflow share exceeds governance threshold",
                investor_risk_framing="Elevated related-party exposure warrants board-approval verification.",
                triggering_transaction_ids=["t4"],
            ),
        ],
        jurisdiction="india",
    )


def _reconciliation_report() -> ReconciliationReport:
    return ReconciliationReport(
        findings=[
            ReconciliationFinding(
                check_name="Total Revenue", claimed_value="Rs.10,00,000", actual_value="Rs.6,20,000",
                direction=MismatchDirection.OVER_REPORTED, delta_pct=61.3, severity="HIGH",
                description="Company claimed Rs.10,00,000 revenue but bank shows Rs.6,20,000.",
                investor_framing="Revenue is overstated by 61% — request source invoices.",
            ),
        ],
        claims_provided=True,
    )


def _customer_analytics() -> CustomerAnalyticsReport:
    return CustomerAnalyticsReport(
        monthly_active_customers={"2024-01": 1, "2024-02": 1},
        monthly_active_trend=0.0,
        churn_events=[
            ChurnEvent(customer_id="cust_ghost", last_payment_date=datetime.date(2023, 10, 1),
                       previous_avg_monthly_spend=Decimal("40000"), months_active_before_churn=5),
        ],
        new_acquisitions={"2024-01": ["cust_alpha"]},
        nrr_per_month={"2024-02": 1.0667},
        cohort_retention={"2024-01": {0: 1.0, 1: 1.0}},
        concentration_trajectory=[
            ConcentrationPoint(month="2024-02", top3_share=1.0, top10_share=1.0),
        ],
        payment_regularity_alerts=[],
        active_customer_ids_by_month={"2024-01": frozenset({"cust_alpha"}), "2024-02": frozenset({"cust_alpha"})},
    )


def _result(**overrides) -> WorkbenchAnalysisResult:
    kwargs = dict(
        doc=_doc(), metrics=_metrics(), risk_report=_risk_report(),
        customer_analytics=_customer_analytics(), compliance_report=_compliance_report(),
        reconciliation_report=_reconciliation_report(), company_name="Acme Pvt Ltd",
    )
    kwargs.update(overrides)
    return WorkbenchAnalysisResult(**kwargs)


def _generate(result: WorkbenchAnalysisResult) -> tuple[BytesIO, bytes]:
    xlsx_bytes, pdf_bytes = WorkbenchLensReport().generate(result)
    return BytesIO(xlsx_bytes), pdf_bytes


def _sheet_text(ws) -> str:
    return " ".join(
        str(ws.cell(r, c).value) for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None
    )


def _pdf_text(pdf_bytes: bytes) -> str:
    with pdfplumber.open(BytesIO(pdf_bytes)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


# ---------------------------------------------------------------------------
# Basic structure
# ---------------------------------------------------------------------------

class TestBasicStructure:
    def test_returns_valid_xlsx_and_pdf(self) -> None:
        xlsx_bytes, pdf_bytes = WorkbenchLensReport().generate(_result())
        assert xlsx_bytes[:2] == b"PK"  # zip magic — xlsx is a zip
        assert pdf_bytes[:5] == b"%PDF-"

    def test_ten_sheets_in_expected_order(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        assert wb.sheetnames == [
            "Executive Summary", "Financial Health", "Risk Flags", "Customer Analytics",
            "Compliance Findings", "Reconciliation", "All Transactions", "Related Parties",
            "Customer Master", "Raw Data Export",
        ]

    def test_pdf_has_four_pages(self) -> None:
        _, pdf_bytes = _generate(_result())
        with pdfplumber.open(BytesIO(pdf_bytes)) as pdf:
            assert len(pdf.pages) == 4


# ---------------------------------------------------------------------------
# Executive Summary — content correctness
# ---------------------------------------------------------------------------

class TestExecutiveSummary:
    def test_kpi_values_present(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Executive Summary"])
        assert "6.20L" in text            # total_revenue = 620000 -> Rs.6.20L
        assert "1.23L" in text            # avg_monthly_burn = 122500 -> Rs.1.23L
        assert "16.3 mo" in text          # runway
        assert "1 HIGH" in text           # compliance high_count and recon high_count both 1
        assert "42/100" not in text       # composite score no longer rendered (Wave 4.5)

    def test_key_findings_lists_risk_flag(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Executive Summary"])
        assert "founder_over_extraction" in text

    def test_flag_summary_table_shows_severity_breakdown(self) -> None:
        """Wave 4.5 — replaces the removed composite risk score."""
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Executive Summary"])
        assert "FLAG SUMMARY" in text
        assert "Risk flags" in text
        assert "Compliance exceptions" in text
        assert "Reconciliation findings" in text

    def test_summary_narrative_shown(self) -> None:
        result = _result()
        result.narrative = "A distinctive workbench summary narrative."
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Executive Summary"])
        assert "SUMMARY" in text
        assert "distinctive workbench summary narrative" in text

    def test_company_name_in_title(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        assert wb["Executive Summary"]["A1"].value == "WORKBENCH ANALYSIS — Acme Pvt Ltd"


# ---------------------------------------------------------------------------
# Financial Health
# ---------------------------------------------------------------------------

class TestFinancialHealth:
    def test_monthly_rows_have_correct_values(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        ws = wb["Financial Health"]
        row2 = [ws.cell(2, c).value for c in range(1, 6)]
        assert row2[0] == "2024-01"
        assert "3.00L" in row2[1]   # revenue
        assert "1.70L" in row2[2]   # burn
        assert "▲ Growing" in row2[4]  # net = 300000-170000 = positive

    def test_february_is_growing_too(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        ws = wb["Financial Health"]
        row3 = [ws.cell(3, c).value for c in range(1, 6)]
        assert row3[0] == "2024-02"
        assert "▲ Growing" in row3[4]  # net = 320000-75000 = positive


# ---------------------------------------------------------------------------
# Risk Flags
# ---------------------------------------------------------------------------

class TestRiskFlagsSheet:
    def test_flag_row_has_correct_detector_and_severity(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Risk Flags"])
        assert "founder_over_extraction" in text
        assert "HIGH" in text
        assert "t4" in text  # triggering transaction id

    def test_total_flags_row_shows_severity_breakdown(self) -> None:
        """Wave 4.5 — replaces the removed composite risk score row."""
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Risk Flags"])
        assert "Total Flags" in text
        assert "1 High / 0 Medium / 0 Low" in text
        assert "42" not in text  # old composite score value must not appear

    def test_no_flags_case(self) -> None:
        result = _result(risk_report=RiskReport(flags=[], composite_score=0.0, narrative="No flags."))
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Risk Flags"])
        assert "founder_over_extraction" not in text


# ---------------------------------------------------------------------------
# Customer Analytics
# ---------------------------------------------------------------------------

class TestCustomerAnalyticsSheet:
    def test_active_customers_and_nrr(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Customer Analytics"])
        assert "107%" in text  # nrr 1.0667 -> 107%

    def test_churned_customer_listed(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Customer Analytics"])
        assert "cust_ghost" in text

    def test_cohort_retention_values(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Customer Analytics"])
        assert "100%" in text  # cohort retention offsets 0 and 1 both 1.0


# ---------------------------------------------------------------------------
# Compliance Findings
# ---------------------------------------------------------------------------

class TestComplianceSheet:
    def test_exception_row_has_correct_fields(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Compliance Findings"])
        assert "IND_RPT_CONCENTRATION" in text
        assert "Companies Act §188" in text
        assert "board-approval verification" in text

    def test_no_exceptions_case(self) -> None:
        result = _result(compliance_report=ComplianceReport(exceptions=[], jurisdiction="india"))
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Compliance Findings"])
        assert "No compliance exceptions found." in text


# ---------------------------------------------------------------------------
# Reconciliation
# ---------------------------------------------------------------------------

class TestReconciliationSheet:
    def test_finding_row_has_correct_values(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Reconciliation"])
        assert "Total Revenue" in text
        assert "Rs.10,00,000" in text
        assert "over_reported" in text
        assert "+61%" in text

    def test_no_claims_provided_case(self) -> None:
        result = _result(reconciliation_report=ReconciliationReport(findings=[], claims_provided=False))
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Reconciliation"])
        assert "No company claims provided" in text

    def test_claims_provided_but_no_findings_case(self) -> None:
        result = _result(reconciliation_report=ReconciliationReport(findings=[], claims_provided=True))
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Reconciliation"])
        assert "reconcile within tolerance" in text


# ---------------------------------------------------------------------------
# All Transactions
# ---------------------------------------------------------------------------

class TestAllTransactionsSheet:
    def test_transaction_values_present(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        ws = wb["All Transactions"]
        rows = [[ws.cell(r, c).value for c in range(1, 10)] for r in range(2, 7)]
        # t1: 2024-01-15, credit 300000, category REVENUE, customer_id cust_alpha
        t1_row = next(r for r in rows if r[0] == "2024-01-15")
        assert t1_row[3] == 300000.0
        assert "REVENUE" in t1_row[5]
        assert t1_row[6] == "cust_alpha"

    def test_related_party_transaction_flagged(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["All Transactions"])
        assert "Founder Family Ventures Pvt Ltd" in text

    def test_anomaly_flags_shown(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["All Transactions"])
        assert "round_amount" in text


# ---------------------------------------------------------------------------
# Related Parties
# ---------------------------------------------------------------------------

class TestRelatedPartiesSheet:
    def test_related_party_row(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Related Parties"])
        assert "Founder Family Ventures Pvt Ltd" in text
        assert "75000" in text or "75,000" in text

    def test_no_related_party_transactions_case(self) -> None:
        doc = _doc()
        # strip the related-party txn
        doc.transactions = [t for t in doc.transactions if not t.is_related_party]
        result = _result(doc=doc)
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Related Parties"])
        assert "No related-party transactions detected." in text


# ---------------------------------------------------------------------------
# Customer Master
# ---------------------------------------------------------------------------

class TestCustomerMasterSheet:
    def test_customer_totals_and_status(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        text = _sheet_text(wb["Customer Master"])
        assert "cust_alpha" in text
        # total revenue for cust_alpha = 300000 + 320000 = 620000 -> Rs.6.20L
        assert "6.20L" in text
        assert "Active" in text  # cust_alpha is in active_customer_ids_by_month for latest month

    def test_status_uses_analyst_computed_active_ids_not_raw_scan(self) -> None:
        """Regression guard for the Wave 3.2 fix: status must come from
        active_customer_ids_by_month, not a fresh scan of doc.transactions."""
        ca = _customer_analytics()
        # Deliberately make the analyst's data say cust_alpha is NOT active in
        # the latest month, even though raw transactions would suggest it is.
        ca.active_customer_ids_by_month = {"2024-01": frozenset(), "2024-02": frozenset()}
        result = _result(customer_analytics=ca)
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        ws = wb["Customer Master"]
        row = [ws.cell(2, c).value for c in range(1, 7)]
        assert row[0] == "cust_alpha"
        assert row[5] == "Inactive"  # not "Active", proving it used the (overridden) analyst data


# ---------------------------------------------------------------------------
# Raw Data Export
# ---------------------------------------------------------------------------

class TestRawDataExportSheet:
    def test_transaction_id_and_anomaly_flags_json(self) -> None:
        buf, _ = _generate(_result())
        wb = load_workbook(buf)
        ws = wb["Raw Data Export"]
        rows = {ws.cell(r, 1).value: r for r in range(2, 7)}
        assert "t5" in rows
        r = rows["t5"]
        flags_json = ws.cell(r, 12).value
        assert json.loads(flags_json) == ["round_amount"]


# ---------------------------------------------------------------------------
# Aggregator caveat (Wave 3.1 wiring, untested until now)
# ---------------------------------------------------------------------------

class TestAggregatorCaveat:
    def test_caveat_in_customer_analytics_sheet_when_dominated(self) -> None:
        ca = _customer_analytics()
        ca.aggregator_revenue_pct = 0.8
        result = _result(customer_analytics=ca)
        buf, _ = _generate(result)
        wb = load_workbook(buf)
        text = _sheet_text(wb["Customer Analytics"])
        assert "aggregator-settled" in text
        assert "80%" in text

    def test_caveat_absent_when_not_dominated(self) -> None:
        buf, _ = _generate(_result())  # default aggregator_revenue_pct = 0.0
        wb = load_workbook(buf)
        text = _sheet_text(wb["Customer Analytics"])
        assert "aggregator-settled" not in text

    def test_caveat_in_pdf_page4(self) -> None:
        ca = _customer_analytics()
        ca.aggregator_revenue_pct = 0.8
        result = _result(customer_analytics=ca)
        _, pdf_bytes = _generate(result)
        assert "aggregator-settled" in _pdf_text(pdf_bytes)


# ---------------------------------------------------------------------------
# PDF content correctness
# ---------------------------------------------------------------------------

class TestPDFContent:
    def test_page1_kpis_and_risk_flag(self) -> None:
        _, pdf_bytes = _generate(_result())
        text = _pdf_text(pdf_bytes)
        assert "Acme Pvt Ltd" in text
        assert "6.20L" in text
        assert "founder_over_extraction" in text

    def test_page2_compliance_citation(self) -> None:
        _, pdf_bytes = _generate(_result())
        text = _pdf_text(pdf_bytes)
        assert "Companies Act §188" in text
        assert "board-approval verification" in text

    def test_page3_reconciliation_framing(self) -> None:
        _, pdf_bytes = _generate(_result())
        text = _pdf_text(pdf_bytes)
        assert "Total Revenue" in text
        assert "overstated by 61%" in text

    def test_page4_churn_customer(self) -> None:
        _, pdf_bytes = _generate(_result())
        text = _pdf_text(pdf_bytes)
        assert "cust_ghost" in text

    def test_no_risk_flags_message(self) -> None:
        result = _result(risk_report=RiskReport(flags=[], composite_score=0.0, narrative="None."))
        _, pdf_bytes = _generate(result)
        assert "No risk flags detected." in _pdf_text(pdf_bytes)

    def test_no_compliance_exceptions_message(self) -> None:
        result = _result(compliance_report=ComplianceReport(exceptions=[], jurisdiction="india"))
        _, pdf_bytes = _generate(result)
        assert "No compliance exceptions identified." in _pdf_text(pdf_bytes)

    def test_no_reconciliation_claims_message(self) -> None:
        result = _result(reconciliation_report=ReconciliationReport(findings=[], claims_provided=False))
        _, pdf_bytes = _generate(result)
        assert "No company claims were provided" in _pdf_text(pdf_bytes)

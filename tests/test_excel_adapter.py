"""Tests for the Excel adapter (Sprint 3, Step 3.6)."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from adapters import AdapterError, ExcelAdapter

_HEADER = ["Date", "Narration", "Chq./Ref.No.", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]
_ROWS = [
    ["01/04/24", "NEFT CR-ACME PVT LTD-INV-1", "REF001", None, 50000.00, 150000.00],
    ["03/04/24", "UPI-VENDOR PAYMENT-XYZ", "REF002", 12000.00, None, 138000.00],
    ["05/04/24", "IMPS-SALARY-STAFF", "REF003", 25000.00, None, 113000.00],
]


def _build_workbook(
    tmp_path: Path,
    header_row: int = 1,
    preamble: list[list[str]] | None = None,
    extra_sheets: dict[str, list[list]] | None = None,
    merge_title: bool = False,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    row_num = 1
    for p in (preamble or []):
        for c, val in enumerate(p, 1):
            ws.cell(row=row_num, column=c, value=val)
        row_num += 1

    if merge_title:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=6)
        ws.cell(row=row_num, column=1, value="Account Statement")
        row_num += 1

    for c, val in enumerate(_HEADER, 1):
        ws.cell(row=row_num, column=c, value=val)
    row_num += 1

    for r in _ROWS:
        for c, val in enumerate(r, 1):
            ws.cell(row=row_num, column=c, value=val)
        row_num += 1

    for name, rows in (extra_sheets or {}).items():
        sheet = wb.create_sheet(name)
        for r_idx, r in enumerate(rows, 1):
            for c_idx, val in enumerate(r, 1):
                sheet.cell(row=r_idx, column=c_idx, value=val)

    path = tmp_path / "stmt.xlsx"
    wb.save(path)
    return path


class TestBasicExtraction:
    def test_extracts_all_transaction_rows(self, tmp_path: Path) -> None:
        path = _build_workbook(tmp_path)
        stmt = ExcelAdapter().extract(path)
        assert len(stmt.rows) == 3

    def test_fields_mapped_correctly(self, tmp_path: Path) -> None:
        path = _build_workbook(tmp_path)
        stmt = ExcelAdapter().extract(path)
        row = stmt.rows[0]
        assert row.date_raw == "01/04/24"
        assert "ACME" in row.description_raw
        assert row.credit_raw in ("50000.0", "50000")
        assert row.debit_raw is None
        assert row.balance_raw in ("150000.0", "150000")


class TestHeaderNotFirstRow:
    def test_header_after_preamble_rows(self, tmp_path: Path) -> None:
        preamble = [["Bank Statement"], ["Account No: 1234567890"], ["Period: Apr 2024"], []]
        path = _build_workbook(tmp_path, preamble=preamble)
        stmt = ExcelAdapter().extract(path)
        assert len(stmt.rows) == 3
        assert stmt.rows[0].date_raw == "01/04/24"


class TestMergedCells:
    def test_merged_title_row_does_not_break_detection(self, tmp_path: Path) -> None:
        path = _build_workbook(tmp_path, merge_title=True)
        stmt = ExcelAdapter().extract(path)
        assert len(stmt.rows) == 3


class TestMultiSheet:
    def test_picks_the_sheet_with_transactions(self, tmp_path: Path) -> None:
        extra = {
            "ReadMe": [["This workbook contains one statement export."], ["Contact support for help."]],
            "Summary": [["Opening Balance", 100000], ["Closing Balance", 113000]],
        }
        path = _build_workbook(tmp_path, extra_sheets=extra)
        stmt = ExcelAdapter().extract(path)
        assert len(stmt.rows) == 3


class TestErrors:
    def test_missing_file_raises(self, tmp_path: Path) -> None:
        with pytest.raises(AdapterError):
            ExcelAdapter().extract(tmp_path / "does_not_exist.xlsx")

    def test_no_detectable_header_raises(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["foo", "bar", "baz"])
        ws.append([1, 2, 3])
        path = tmp_path / "junk.xlsx"
        wb.save(path)
        with pytest.raises(AdapterError):
            ExcelAdapter().extract(path)
